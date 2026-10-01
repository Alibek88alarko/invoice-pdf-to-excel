import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
import invoice_to_excel as ite  # noqa: E402
import make_samples  # noqa: E402


@pytest.fixture(scope="module")
def results(tmp_path_factory):
    make_samples.OUT = tmp_path_factory.mktemp("samples")
    for inv in make_samples.INVOICES:
        make_samples.draw(inv)
    return {r["file"]: r for r in ite.process(make_samples.OUT)}


def test_classic_layout_fields(results):
    r = results["northwind_INV-1042.pdf"]
    assert r["supplier"] == "Northwind Office Supply Ltd"
    assert r["number"] == "INV-1042" and r["date"] == "2026-09-12" and r["currency"] == "GBP"
    assert len(r["lines"]) == 3
    assert r["subtotal"] == 262.47 and r["total"] == 314.96
    assert r["status"] == "OK"


def test_compact_layout_and_date_normalised(results):
    r = results["bluepeak_2026-0377.pdf"]
    assert r["number"] == "2026-0377" and r["date"] == "2026-09-18" and r["currency"] == "EUR"
    assert r["status"] == "OK"


def test_mismatch_goes_to_check(results):
    r = results["greenline_GL-88.pdf"]
    assert r["status"] == "CHECK"
    assert any("total says 590.00" in p for p in r["problems"])


def test_unreadable_file_is_flagged_not_dropped(tmp_path):
    (tmp_path / "broken.pdf").write_bytes(b"not a pdf")
    out = ite.process(tmp_path)
    assert out[0]["status"] == "CHECK" and "could not read" in out[0]["problems"][0]


def test_excel_has_three_sheets(results, tmp_path):
    out = tmp_path / "inv.xlsx"
    ite.write_excel(list(results.values()), out)
    wb = load_workbook(out)
    assert wb.sheetnames == ["Invoices", "Lines", "Check"]
    assert wb["Lines"].max_row == 1 + 3 + 2 + 2
    assert wb["Check"].max_row >= 2
